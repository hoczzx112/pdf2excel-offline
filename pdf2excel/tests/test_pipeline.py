"""Kiểm thử end-to-end cho PDF → Excel.

Chạy được theo 2 cách:
    python tests/test_pipeline.py        # không cần cài pytest
    pytest tests/ -q                     # nếu có pytest

Test dùng PDF tổng hợp do tests/make_pdf.py sinh ra (không cần file mẫu,
không cần mạng).
"""
from __future__ import annotations

import io
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from make_pdf import build_table_pdf  # noqa: E402

from app.core.extract import convert  # noqa: E402
from app.export.xlsx import _clean, save_workbook  # noqa: E402


def _tmpdir() -> Path:
    return Path(tempfile.mkdtemp(prefix="pdf2xlsx_test_"))


# --------------------------------------------------------------------- unit
def test_clean_xml_chars() -> None:
    """Ký tự không hợp lệ với XML phải bị loại (lỗi thật gặp trên file WI)."""
    raw = "外圈滚道直径\ufffe目测\n\x00\x07giữ\t tab"
    out = _clean(raw)
    assert "\ufffe" not in out
    assert "\x00" not in out and "\x07" not in out
    assert "\n" in out and "\t" in out
    assert out.startswith("外圈滚道直径")


# ------------------------------------------------------- hình học ô / dấu "/"
def test_looks_like_slash() -> None:
    """PP-OCR đọc "/" thành "1": phân biệt bằng độ thẳng của nét mực."""
    import numpy as np

    from app.core.ocr import looks_like_slash

    h = w = 32
    # dấu "/" quét 200 DPI: một nét chéo đều, dx/dy ~ 0.24
    slash = np.zeros((h, w), np.uint8)
    for y in range(3, 29):
        x = int(round(6 + (y - 3) * 0.24))
        slash[y, x: x + 2] = 255

    # chữ "1": nét đứng + gạch chân + móc đầu -> không thẳng hàng
    one = np.zeros((h, w), np.uint8)
    one[4:28, 15:18] = 255
    one[25:28, 8:23] = 255
    one[4:7, 11:16] = 255

    assert looks_like_slash(slash) is True
    assert looks_like_slash(one) is False


def test_leak_region_dropped() -> None:
    """Vùng lớn bị "rò rỉ" (nuốt nhiều ô thật) phải bị bỏ, không nuốt bảng."""
    from app.core.cellcomponents import RawCell, _prune_nested

    leak = RawCell(0, 0, 1000, 1000)
    real = [RawCell(100, 100, 400, 400), RawCell(100, 420, 400, 700)]
    kept = _prune_nested([leak] + real)
    assert len(kept) == 2, kept
    assert [(c.x0, c.y0) for c in kept] == [(100, 100), (100, 420)], kept

    # Ô lớn chỉ chứa mảnh vụn (bản vẽ/logo): giữ ô lớn, bỏ mảnh vụn.
    big = RawCell(0, 0, 1000, 1000)
    frags = [RawCell(100, 100, 130, 130), RawCell(200, 200, 230, 230)]
    assert _prune_nested([big] + frags) == [big]


def test_build_grid_keeps_real_cells_and_drops_junk() -> None:
    """Lưới phải giữ đủ ô thật và loại mảnh vụn 13x13 trong logo."""
    from app.core.cellcomponents import RawCell, build_grid

    xs = [100, 300, 500, 700]
    ys = [100, 200, 300]
    cells = [
        RawCell(xs[i], ys[j], xs[i + 1], ys[j + 1])
        for i in range(3)
        for j in range(2)
    ]
    junk = RawCell(880, 800, 893, 813)
    rows, cols, idx = build_grid(cells + [junk])
    assert len(idx) == 6, (len(rows), len(cols), len(idx))
    assert all(c.x0 < 800 for c in cells), "ô thật không nằm trong vùng rác"


def test_fill_grid_holes_restores_missing_column() -> None:
    """Cột bị "rò" mất nhưng có chữ + đủ nét kẻ thì phải được dựng lại."""
    import numpy as np

    from app.core.cellcomponents import _fill_grid_holes
    from app.core.gridlines import CellRect
    from app.core.models import TextItem

    rows = [0, 100, 200, 300]
    cols = [0, 100, 200, 300]
    # nét kẻ: mọi đường bao quanh đều có (hàng 1..2, cột 1..2)
    hl = np.zeros((320, 320), np.uint8)
    for y in (0, 100, 200, 300):
        hl[max(0, y - 2): y + 3, :] = 255
    vl = np.zeros((320, 320), np.uint8)
    for x in (0, 100, 200, 300):
        vl[:, max(0, x - 2): x + 3] = 255

    # ô (1,0) bị thiếu; các ô còn lại đã dò được
    have = [
        CellRect(1, 1, 1, 2, 100, 100, 300, 200),
        CellRect(2, 0, 2, 2, 0, 200, 300, 300),
    ]
    items = [
        TextItem(x0=10, y0=110, x1=90, y1=190, text="1", score=1.0),
        TextItem(x0=110, y0=110, x1=190, y1=190, text="x", score=1.0),
    ]
    out = _fill_grid_holes(have, rows, cols, hl, vl, items)
    assert len(out) == 3, out
    added = [c for c in out if c not in have][0]
    assert (added.r0, added.c0) == (1, 0), added


def test_superscript_keeps_line_order() -> None:
    """Chỉ số trên (10⁸) không được tách thành dòng riêng làm đảo thứ tự chữ.

    Lỗi thật trên file WI: ô "轴向刚度(0游隙)：1.51×10⁸ N/m MIN" bị ghép thành
    "隙1 51×108" + ")：." vì số 8 nằm cao hơn dòng.
    """
    from app.core.models import TextItem
    from app.core.textlayer import assemble_text

    def ch(x, y, t, h=20, dh=8):
        return TextItem(x0=x, y0=y, x1=x + 12, y1=y + h, text=t)

    items = [
        ch(0, 0, "轴"), ch(13, 0, "向"), ch(26, 0, "刚"), ch(39, 0, "度"),
        ch(52, 1, "("), ch(65, 2, "0"), ch(78, 0, "游"),
        # dòng 2, số 8 là chỉ số trên nên cao hơn 6 px và nhỏ hơn
        ch(0, 24, "隙"), ch(13, 26, ")"), ch(26, 27, "："),
        ch(39, 26, "1"), ch(52, 26, "."), ch(65, 26, "5"),
        ch(78, 26, "1"), ch(91, 26, "×"), ch(104, 26, "1"),
        ch(117, 26, "0"), ch(130, 22, "8", h=12),
        # dòng 3
        ch(0, 48, "N"), ch(13, 49, "/"), ch(26, 49, "m"),
    ]
    text, _ = assemble_text(items, -5, -5, 200, 90)
    assert text.split("\n")[:2] == ["轴向刚度(0游", "隙)：1.51×108"], text


def test_text_erase_keeps_touching_grid_line() -> None:
    """Hộp chữ chạm sát đường kẻ không được xoá mất đường kẻ (làm gộp 2 ô)."""
    import numpy as np

    from app.core.gridlines import remove_text_pixels
    from app.core.models import TextItem

    binary = np.zeros((80, 80), np.uint8)
    binary[:, 49:52] = 255          # đường kẻ dọc 3 px
    # hộp chữ kết thúc đúng ở mép đường kẻ (như pdfium trả về)
    item = TextItem(x0=10, y0=20, x1=49, y1=45, text="色")
    clean = remove_text_pixels(binary, [item])
    assert int((clean[:, 50] > 0).sum()) == 80, "đường kẻ bị xoá mất"
    assert int((clean[20:45, 10:49] > 0).sum()) == 0, "vùng chữ phải được xoá"



# ---------------------------------------------------------------- pipeline
def test_convert_synthetic_table() -> None:
    tmp = _tmpdir()
    pdf = build_table_pdf(tmp / "syn.pdf")
    res = convert(pdf)

    assert len(res.pages) == 1, res.pages
    assert res.pages[0].kind == "text"
    assert len(res.tables) == 1, [t.id for t in res.tables]

    table = res.tables[0]
    assert (table.num_rows, table.num_cols) == (3, 3)
    grid = table.grid()
    assert grid[0] == ["Code", "Name", "Qty"]
    assert grid[1] == ["A1", "Bolt M6", "10"]
    assert grid[2] == ["A2", "Nut M6", "20"]


def test_workbook_structure() -> None:
    from openpyxl import load_workbook

    tmp = _tmpdir()
    res = convert(build_table_pdf(tmp / "syn.pdf"))
    out = save_workbook(res, tmp / "out.xlsx")

    wb = load_workbook(out)
    assert wb.sheetnames == ["T1-B1"], wb.sheetnames

    ws = wb["T1-B1"]
    assert ws["A1"].value == "Code"
    assert ws["B2"].value == "Bolt M6"
    assert ws["A1"].font.bold is True
    assert ws.max_row == 3 and ws.max_column == 3


def test_soft_join_removes_hard_line_breaks() -> None:
    """Ô dài phải để Excel tự xuống dòng, không chèn "\\n" (Alt+Enter)."""
    from app.export.xlsx import _soft_join

    assert _soft_join("轴向刚度(0游\n隙)：1.51×108") == "轴向刚度(0游隙)：1.51×108"
    assert _soft_join("GM08油脂\n(Mobil Polyrex\nEM)") == "GM08油脂(Mobil Polyrex EM)"
    assert _soft_join("CRB11116508\nCX3.01") == "CRB11116508 CX3.01"
    assert _soft_join("a\n\nb") == "a b"
    assert _soft_join("") == ""


def test_comma_stays_with_digits() -> None:
    """Dấu phẩy thấp hơn chữ số nhưng vẫn phải nằm cùng dòng: "1,2"."""
    from app.core.models import TextItem
    from app.core.textlayer import assemble_text

    # đúng toạ độ thật của file WI-ZDDJ (dấu phẩy thấp hơn và lệch phải)
    items = [
        TextItem(x0=1425.4, y0=428.5, x1=1431.8, y1=448.9, text="1"),
        TextItem(x0=1441.6, y0=428.5, x1=1453.2, y1=448.9, text="2"),
        TextItem(x0=1436.8, y0=446.1, x1=1439.3, y1=452.9, text=","),
    ]
    text, _ = assemble_text(items, 1399, 420, 1477, 498)
    assert text == "1,2", text

    # bản PDF khác dùng dấu phẩy fullwidth ở cùng ô -> chuẩn hoá về ","
    wide = [
        TextItem(x0=1416.7, y0=702.9, x1=1423.1, y1=723.3, text="1"),
        TextItem(x0=1450.4, y0=702.9, x1=1461.9, y1=723.3, text="2"),
        TextItem(x0=1428.6, y0=716.7, x1=1432.9, y1=724.4, text="，"),
    ]
    text, _ = assemble_text(wide, 1399, 694, 1477, 772)
    assert text == "1,2", text


def test_no_meta_sheets_in_workbook() -> None:
    """File Excel chỉ có sheet bảng, không có _Thông tin / _Văn bản gốc."""
    from openpyxl import load_workbook

    tmp = _tmpdir()
    res = convert(build_table_pdf(tmp / "only.pdf"))
    out = save_workbook(res, tmp / "only.xlsx")
    names = load_workbook(out).sheetnames
    assert names == ["T1-B1"], names
    assert all(not n.startswith("_") for n in names)


def test_junction_gap_between_grid_lines_is_sealed() -> None:
    """Đường kẻ ngang dừng cách đường dọc 1 px làm 2 ô dính nhau.

    Lỗi thật trên WI-TSG trang 3: ô "60s" nuốt luôn ô "300s" vì góc chữ T hở
    đúng 1 px, vùng trắng chảy qua khe.
    """
    import numpy as np

    from app.core.cellcomponents import _seal_line_gaps, find_cell_interiors

    b = np.zeros((140, 140), np.uint8)
    b[0:3, :] = 255          # khung trên
    b[-3:, :] = 255          # khung dưới
    b[:, 0:3] = 255          # khung trái
    b[:, -3:] = 255          # khung phải
    b[:, 68:71] = 255        # đường dọc chia 2 cột
    b[68:71, 3:67] = 255     # đường ngang nửa trái, dừng sớm 1 px -> khe ở x=67

    def left_cells(binary):
        return [
            c for c in find_cell_interiors(binary)
            if c.x1 <= 70 and c.y0 > 0 and c.y1 < 139
        ]

    assert len(left_cells(b)) == 1, "chưa vá thì 2 ô phải dính thành 1"

    sealed = _seal_line_gaps(b)
    assert int((sealed[69, 67] > 0)) == 1, "khe ở góc chữ T phải được vá"
    assert len(left_cells(sealed)) == 2, left_cells(sealed)


def test_formula_text_not_treated_as_formula() -> None:
    """Ô nội dung bắt đầu bằng '=' phải lưu dạng chuỗi, không thành công thức."""
    from openpyxl import load_workbook

    tmp = _tmpdir()
    pdf = build_table_pdf(
        tmp / "eq.pdf",
        header=("A", "B", "C"),
        rows=(("=SUM(1,2)", "x", "y"),),
    )
    res = convert(pdf)
    out = save_workbook(res, tmp / "eq.xlsx")
    wb = load_workbook(out)
    cell = wb["T1-B1"]["A2"]
    assert cell.value == "=SUM(1,2)"
    assert cell.data_type == "s"


# --------------------------------------------------------------------- API
def test_http_flow() -> None:
    """upload → convert → poll → export → download → delete."""
    from fastapi.testclient import TestClient

    from app.main import app

    tmp = _tmpdir()
    pdf = build_table_pdf(tmp / "api.pdf")
    client = TestClient(app)

    assert client.get("/api/health").json()["ok"] is True

    with open(pdf, "rb") as fh:
        resp = client.post(
            "/api/jobs",
            files={"files": ("api.pdf", fh, "application/pdf")},
        )
    assert resp.status_code == 200, resp.text
    job = resp.json()["jobs"][0]
    assert job["analysis"]["num_pages"] == 1
    job_id = job["id"]

    try:
        assert client.post(
            f"/api/jobs/{job_id}/convert",
            json={"mode": "auto", "engine": "fast"},
        ).json()["ok"] is True

        deadline = time.time() + 120
        snap: dict = {}
        while time.time() < deadline:
            snap = client.get(f"/api/jobs/{job_id}").json()
            if snap["status"] in ("done", "error"):
                break
            time.sleep(0.2)
        assert snap["status"] == "done", snap.get("error")
        assert snap["result"]["num_tables"] == 1

        # tải trực tiếp (chưa gọi export) — server phải tự xuất
        auto = client.get(f"/api/jobs/{job_id}/download")
        assert auto.status_code == 200, auto.text
        assert auto.content[:2] == b"PK"

        assert client.post(
            f"/api/jobs/{job_id}/export", json={"include": None, "edits": {}}
        ).status_code == 200

        dl = client.get(f"/api/jobs/{job_id}/download")
        assert dl.status_code == 200
        assert dl.content[:2] == b"PK"          # file .xlsx = zip
        assert len(dl.content) > 4000

        png = client.get(f"/api/jobs/{job_id}/page/1")
        assert png.status_code == 200
        assert png.content[:8] == b"\x89PNG\r\n\x1a\n"
    finally:
        assert client.delete(f"/api/jobs/{job_id}").status_code == 200

    assert client.get(f"/api/jobs/{job_id}").status_code == 404


def test_rejects_non_pdf() -> None:
    from fastapi.testclient import TestClient

    from app.main import app

    client = TestClient(app)
    resp = client.post(
        "/api/jobs",
        files={"files": ("note.txt", io.BytesIO(b"hello"), "text/plain")},
    )
    assert resp.status_code == 400
    assert "PDF" in resp.json()["detail"]


# -------------------------------------------------------------------- runner
def main() -> int:
    tests = [
        (name, obj)
        for name, obj in sorted(globals().items())
        if name.startswith("test_") and callable(obj)
    ]
    failed = 0
    for name, fn in tests:
        t0 = time.time()
        try:
            fn()
            print(f"  PASS  {name}  ({time.time() - t0:.2f}s)")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"  FAIL  {name}: {type(exc).__name__}: {exc}")
    print(f"\n{len(tests) - failed}/{len(tests)} test đạt")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
