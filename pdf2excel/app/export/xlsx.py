"""Xuất kết quả ra Excel (.xlsx) bằng openpyxl."""
from __future__ import annotations

import math
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..config import OCR_WARN_SCORE
from ..core.models import DocResult, Table
from ..core.textlayer import is_cjk

# Ký tự không hợp lệ với XML: C0 control (trừ \t \n), C1, surrogate lẻ,
# và noncharacter (U+FFFE/U+FFFF, U+FDD0..U+FDEF). PDF hay map ký hiệu
# (ví dụ "Φ") sang U+FFFE, openpyxl sẽ báo lỗi nếu không loại bỏ.
_ILLEGAL = re.compile(
    "[\u0000-\u0008\u000b\u000c\u000e-\u001f\u007f-\u009f"
    "\ud800-\udfff\ufdd0-\ufdef\ufffe\uffff]"
)

_THIN = Side(style="thin", color="B0B0B0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FILL = PatternFill("solid", fgColor="E8EEF7")
_WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
_HEAD_FONT = Font(bold=True)

_MAX_COL_WIDTH = 60
_MIN_COL_WIDTH = 8
_MIN_ROW_HEIGHT = 18.0
_MAX_ROW_HEIGHT = 240.0
_ROW_UNIT = 15.0


def _clean(text: str) -> str:
    text = _ILLEGAL.sub("", text or "")
    # noncharacter ở các mặt phẳng phụ (U+1FFFE, U+2FFFF, ...)
    if any((ord(ch) & 0xFFFE) == 0xFFFE for ch in text):
        text = "".join(ch for ch in text if (ord(ch) & 0xFFFE) != 0xFFFE)
    return text


def _soft_join(text: str) -> str:
    """Nối các dòng của một ô thành một đoạn, KHÔNG chèn xuống dòng cứng.

    Trước đây mỗi dòng trong ô được ghi kèm "\\n" (tương đương Alt+Enter), nên
    Excel hiển thị đúng chỗ ngắt của bản PDF nhưng không bao giờ tự dàn lại
    theo độ rộng cột. Nay chỉ nối lại thành câu và để `wrap_text` tự xuống
    dòng. Giữa hai chữ Latin cần một khoảng trắng (PDF tách dòng giữa từ), còn
    cạnh chữ Hán/CJK thì nối liền.
    """
    parts = [p.strip() for p in (text or "").split("\n")]
    parts = [p for p in parts if p]
    if not parts:
        return ""
    out = parts[0]
    for part in parts[1:]:
        if is_cjk(out[-1:]) or is_cjk(part[:1]):
            out += part
        else:
            out += " " + part
    return out


def _set_text(cell, text: str) -> None:
    cell.value = _clean(text)
    # tránh openpyxl hiểu nhầm nội dung bắt đầu bằng "=" thành công thức
    if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"


def _display_width(s: str) -> int:
    """Độ rộng hiển thị (ký tự CJK tính bằng 2)."""
    w = 0
    for ch in s:
        w += 2 if ord(ch) > 0x2E7F else 1
    return w


def _sheet_title(table: Table, used: set[str]) -> str:
    base = f"T{table.page}-B{table.index}"
    title = base
    n = 2
    while title in used:
        title = f"{base}_{n}"
        n += 1
    used.add(title)
    return title[:31]


def _write_table(ws: Worksheet, table: Table, overrides: dict[str, str]) -> None:
    grid = table.grid()
    spans = {(c.row, c.col): (c.rowspan, c.colspan) for c in table.cells}
    scores = {(c.row, c.col): c.score for c in table.cells}

    for key, val in (overrides or {}).items():
        try:
            r_s, c_s = key.split(",")
            r, c = int(r_s), int(c_s)
        except Exception:
            continue
        if 0 <= r < table.num_rows and 0 <= c < table.num_cols:
            grid[r][c] = val

    for (r, c), (rs, cs) in spans.items():
        cell = ws.cell(row=r + 1, column=c + 1)
        raw = grid[r][c] if r < len(grid) and c < len(grid[r]) else ""
        _set_text(cell, _soft_join(str(raw)))
        cell.alignment = Alignment(
            wrap_text=True, vertical="center", horizontal="left"
        )
        score = scores.get((r, c))
        if score is not None and score < OCR_WARN_SCORE:
            cell.fill = _WARN_FILL
            cell.comment = Comment(f"Độ tin cậy OCR thấp: {score:.2f}", "pdf2excel")

    if table.num_rows >= 1:
        for c in range(table.num_cols):
            cell = ws.cell(row=1, column=c + 1)
            cell.font = _HEAD_FONT
            cell.fill = _HEAD_FILL

    for r in range(table.num_rows):
        for c in range(table.num_cols):
            ws.cell(row=r + 1, column=c + 1).border = _BORDER

    for (r, c), (rs, cs) in spans.items():
        if rs > 1 or cs > 1:
            ws.merge_cells(
                start_row=r + 1, start_column=c + 1,
                end_row=r + rs, end_column=c + cs,
            )

    # độ rộng cột: lấy theo dòng dài nhất của bản gốc (trước khi nối), chia đều
    # cho số cột mà ô gộp chiếm giữ.
    col_width = [_MIN_COL_WIDTH] * table.num_cols
    for (r, cc), (rs, cs) in spans.items():
        val = grid[r][cc] if r < len(grid) and cc < len(grid[r]) else ""
        longest = max((_display_width(x) for x in str(val).split("\n")), default=0)
        share = int(longest / max(1, cs)) + 3
        for k in range(cc, min(table.num_cols, cc + cs)):
            col_width[k] = max(col_width[k], min(_MAX_COL_WIDTH, share))
    for c in range(table.num_cols):
        ws.column_dimensions[get_column_letter(c + 1)].width = col_width[c]

    # chiều cao dòng: ước lượng số dòng cần khi chữ tự xuống dòng theo độ rộng
    # cột đã đặt. Ô gộp nhiều hàng thì chia đều cho các hàng đó.
    for r in range(table.num_rows):
        need = 1.0
        for (rr, cc), (rs, cs) in spans.items():
            if rr != r:
                continue
            val = grid[r][cc] if r < len(grid) and cc < len(grid[r]) else ""
            text = _soft_join(str(val))
            avail = max(2.0, sum(col_width[cc:cc + cs]) - 1.5)
            lines = max(1, math.ceil(_display_width(text) / avail))
            if scores.get((rr, cc)) is not None and (scores[(rr, cc)] or 1) < OCR_WARN_SCORE:
                lines += 1
            need = max(need, lines / max(1, rs))
        ws.row_dimensions[r + 1].height = max(
            _MIN_ROW_HEIGHT, min(_MAX_ROW_HEIGHT, _ROW_UNIT * math.ceil(need) + 5)
        )

    ws.freeze_panes = "A2"


def build_workbook(
    result: DocResult,
    include: set[str] | None = None,
    edits: dict[str, dict[str, str]] | None = None,
) -> Workbook:
    """Chỉ ghi các sheet bảng. Không tạo sheet phụ (_Thông tin / _Văn bản gốc)."""
    wb = Workbook()
    default = wb.active
    default.title = "Kết quả"

    used: set[str] = set()
    first = True
    for table in result.tables:
        if include is not None and table.id not in include:
            continue
        if first:
            ws = default
            ws.title = _sheet_title(table, used)
            first = False
        else:
            ws = wb.create_sheet(_sheet_title(table, used))
        _write_table(ws, table, (edits or {}).get(table.id, {}))

    if first:
        # không có bảng nào: giữ một sheet trống để file Excel vẫn hợp lệ
        ws = default
        ws.title = "Kết quả"
        ws["A1"] = "Không có bảng nào được chọn để xuất."
        ws.column_dimensions["A"].width = 50

    return wb


def save_workbook(
    result: DocResult,
    path: Path,
    include: set[str] | None = None,
    edits: dict[str, dict[str, str]] | None = None,
) -> Path:
    wb = build_workbook(result, include=include, edits=edits)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
